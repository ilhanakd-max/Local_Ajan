import sys
import os
import json
import tempfile
from pathlib import Path

sys.path.insert(0, "src")

from lokal_ajan.tools.registry import registry
from lokal_ajan.tools.office_tools import ReadExcelTool, WriteExcelTool, ReadPdfTool, WritePdfTool, AppendExcelTool, EditExcelCellsTool
from lokal_ajan.tools.fs_tools import ReadFileTool, WriteFileTool

def test_office_tools_registration():
    schemas = registry.get_all_schemas()
    tool_names = [t["name"] for t in schemas]
    assert "read_excel" in tool_names
    assert "write_excel" in tool_names
    assert "append_excel" in tool_names
    assert "edit_excel_cells" in tool_names
    assert "read_pdf" in tool_names
    assert "write_pdf" in tool_names

def test_excel_write_and_read():
    with tempfile.TemporaryDirectory() as tmpdir:
        w_excel = WriteExcelTool(workdir=tmpdir)
        r_excel = ReadExcelTool(workdir=tmpdir)
        excel_path = os.path.join(tmpdir, "test.xlsx")

        data = [
            {"Ders": "Matematik", "Saat": "09:00", "Öğretmen": "Ahmet"},
            {"Ders": "Fizik", "Saat": "10:30", "Öğretmen": "Ayşe"}
        ]
        
        # Write
        res_write = w_excel.run(path=excel_path, data_json=json.dumps(data))
        assert "Successfully wrote" in res_write
        assert os.path.exists(excel_path)

        # Read normal
        res_read = r_excel.run(path=excel_path)
        assert "Matematik" in res_read
        assert "Fizik" in res_read

        # Read with max_rows
        res_read_limit = r_excel.run(path=excel_path, max_rows=1)
        assert "Matematik" in res_read_limit
        assert "ilk 1 satır gösterildi" in res_read_limit

def test_pdf_write_and_read():
    with tempfile.TemporaryDirectory() as tmpdir:
        w_pdf = WritePdfTool(workdir=tmpdir)
        r_pdf = ReadPdfTool(workdir=tmpdir)
        pdf_path = os.path.join(tmpdir, "test.pdf")

        content = "Çakabey Kültür Merkezi\n2026-2027 Ders Programı\nPazartesi: Matematik"

        # Write
        res_write = w_pdf.run(path=pdf_path, content=content)
        assert "Successfully wrote" in res_write
        assert os.path.exists(pdf_path)

        # Read
        res_read = r_pdf.run(path=pdf_path)
        assert "Çakabey" in res_read or "Matematik" in res_read

def test_fs_tools_binary_hints():
    with tempfile.TemporaryDirectory() as tmpdir:
        r_tool = ReadFileTool(workdir=tmpdir)
        w_tool = WriteFileTool(workdir=tmpdir)

        pdf_path = os.path.join(tmpdir, "document.pdf")
        xlsx_path = os.path.join(tmpdir, "sheet.xlsx")

        # read_file hints
        read_pdf_err = r_tool.run(path=pdf_path)
        assert "Cannot read PDF with 'read_file'" in read_pdf_err
        assert "'read_pdf'" in read_pdf_err

        read_xlsx_err = r_tool.run(path=xlsx_path)
        assert "Cannot read Excel with 'read_file'" in read_xlsx_err
        assert "'read_excel'" in read_xlsx_err

        # write_file hints
        write_pdf_err = w_tool.run(path=pdf_path, content="test")
        assert "Cannot write PDF with 'write_file'" in write_pdf_err
        assert "'write_pdf'" in write_pdf_err

        write_xlsx_err = w_tool.run(path=xlsx_path, content="test")
        assert "Cannot write Excel with 'write_file'" in write_xlsx_err
        assert "'write_excel'" in write_xlsx_err

def test_excel_append_and_edit_preserve_styles_and_formulas():
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    with tempfile.TemporaryDirectory() as tmpdir:
        excel_path = os.path.join(tmpdir, "styled.xlsx")
        
        # 1. Create a styled template with formula
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Item", "Qty", "Price", "Total"])
        
        # Header style
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

        # Row 2 with formula and custom font
        ws.append(["Pen", 10, 5, "=B2*C2"])
        ws["A2"].font = Font(italic=True, color="FF0000")
        ws["A2"].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        wb.save(excel_path)

        # 2. Append new row using AppendExcelTool
        appender = AppendExcelTool(workdir=tmpdir)
        res_append = appender.run(
            path=excel_path,
            rows_json=json.dumps([["Notebook", 20, 15, "=B3*C3"]]),
            copy_style=True
        )
        assert "Successfully appended" in res_append

        # Verify appended row and preserved styles
        wb_check = openpyxl.load_workbook(excel_path, data_only=False)
        ws_check = wb_check["Sales"]
        assert ws_check.max_row == 3
        assert ws_check["A3"].value == "Notebook"
        assert ws_check["D2"].value == "=B2*C2"  # Formula preserved!
        assert ws_check["D3"].value == "=B3*C3"  # New formula set!
        assert ws_check["A3"].font.italic is True  # Style copied from A2!
        assert ws_check["A3"].font.color.rgb == "00FF0000" or ws_check["A3"].font.color.rgb == "FF0000"

        # 3. Edit cell using EditExcelCellsTool
        editor = EditExcelCellsTool(workdir=tmpdir)
        res_edit = editor.run(
            path=excel_path,
            updates_json=json.dumps({"B2": 12, "E1": "Status"})
        )
        assert "Successfully updated" in res_edit

        wb_check2 = openpyxl.load_workbook(excel_path, data_only=False)
        ws_check2 = wb_check2["Sales"]
        assert ws_check2["B2"].value == 12
        assert ws_check2["D2"].value == "=B2*C2"  # Formula still preserved!
        assert ws_check2["E1"].value == "Status"

